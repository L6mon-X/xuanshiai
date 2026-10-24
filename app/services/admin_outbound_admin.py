"""电话外呼（A6）服务：平台配置、余额、外显号码池、外呼名单、通话记录。

真实外呼由供应商适配层执行（``app/services/outbound_call``）。未配置供应商时
所有需要外呼的动作在写库之前就 503，不会留下「calling」僵尸名单行。
"""

from __future__ import annotations

import uuid
from datetime import datetime
from io import BytesIO
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.sensitive_fields import mask_phone
from app.schemas.admin_outbound_admin import (
    OutboundBalanceItem,
    OutboundConfig,
    OutboundConfigUpdate,
    OutboundNumberCreate,
    OutboundNumberItem,
    OutboundNumberPage,
    OutboundNumberUpdate,
    OutboundRecordItem,
    OutboundRecordPage,
    OutboundTaskImportResult,
    OutboundTaskItem,
    OutboundTaskPage,
)
from app.services import admin_config as config_service
from app.services.outbound_call.providers import get_outbound_provider, provider_status

_MAX_IMPORT_ROWS = 2000
_NAMESPACE = "sys_outbound"

_MATCHMAKER_VALID_SQL = """SELECT 1 FROM users u WHERE u.id = :mid AND u.status = 1
    AND EXISTS (SELECT 1 FROM user_role ur WHERE ur.user_id = u.id
        AND ur.role_code = 'service_matchmaker' AND ur.status = 1)"""

_NUMBER_SELECT = """SELECT n.id, n.phone, n.area, n.status, n.bound_matchmaker_id, n.remark, n.created_at,
    u.nickname AS bound_matchmaker_name
    FROM admin_outbound_number n LEFT JOIN users u ON u.id = n.bound_matchmaker_id"""

_TASK_SELECT = """SELECT t.id, t.batch_no, t.member_id, t.name, t.phone, t.matchmaker_id,
    u.nickname AS matchmaker_name, t.status, t.result, t.remark, t.call_count,
    t.last_called_at, t.created_at
    FROM admin_outbound_task t LEFT JOIN users u ON u.id = t.matchmaker_id"""

_RECORD_SELECT = """SELECT r.id, r.task_id, r.phone, r.matchmaker_id, u.nickname AS matchmaker_name,
    r.direction, r.result, r.duration, r.recording_url, r.provider_call_id,
    r.started_at, r.ended_at, r.created_at
    FROM admin_outbound_record r LEFT JOIN users u ON u.id = r.matchmaker_id"""


# ─────────────────────────── 配置与余额 ───────────────────────────


async def get_config(db: AsyncSession) -> OutboundConfig:
    snapshot = await config_service.get_config(db, _NAMESPACE)
    config = snapshot.config if isinstance(snapshot.config, dict) else {}
    status = provider_status()
    return OutboundConfig(
        provider=str(status["provider"]),
        configured=bool(status["configured"]),
        base_url_configured=bool(status["base_url_configured"]),
        account_configured=bool(status["account_configured"]),
        secret_configured=bool(status["secret_configured"]),
        account_name=config.get("account_name"),
        call_center_url=config.get("call_center_url"),
        record_download_url=config.get("record_download_url"),
        version=snapshot.version,
        updated_at=snapshot.updated_at,
    )


async def update_config(
    db: AsyncSession, admin_id: int, body: OutboundConfigUpdate
) -> OutboundConfig:
    changes = body.model_dump(exclude_unset=True)
    await config_service.patch_config(
        db, admin_id, _NAMESPACE, changes, summary=f"更新外呼平台配置：{'/'.join(changes)}"
    )
    return await get_config(db)


async def get_balance(db: AsyncSession) -> OutboundBalanceItem:
    """余额以供应商实时查询为准；失败不落库，避免把异常值当余额展示。"""
    provider = get_outbound_provider()
    try:
        balance = await provider.fetch_balance()
    except HTTPException:
        raise
    except Exception as exc:  # 供应商网络异常
        raise HTTPException(502, detail="外呼平台余额查询失败") from exc
    checked_at = datetime.utcnow()
    await db.execute(
        text(
            "INSERT INTO admin_outbound_account (provider, account_name, balance, currency, checked_at)"
            " VALUES (:provider, :account_name, :balance, :currency, :checked_at)"
        ),
        {
            "provider": balance.provider,
            "account_name": balance.account_name,
            "balance": balance.balance,
            "currency": balance.currency,
            "checked_at": checked_at,
        },
    )
    await db.commit()
    return OutboundBalanceItem(
        provider=balance.provider,
        account_name=balance.account_name,
        balance=float(balance.balance),
        currency=balance.currency,
        checked_at=checked_at,
    )


# ─────────────────────────── 外显号码池 ───────────────────────────


def _number_item(row: Any) -> OutboundNumberItem:
    data = dict(row)
    data["phone"] = mask_phone(data.get("phone")) or ""
    return OutboundNumberItem(**data)


async def _get_number(db: AsyncSession, number_id: int) -> OutboundNumberItem:
    row = (
        await db.execute(text(f"{_NUMBER_SELECT} WHERE n.id = :id"), {"id": number_id})
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="外显号码不存在")
    return _number_item(row)


async def list_numbers(
    db: AsyncSession, page: int, page_size: int, status: int | None, keyword: str | None
) -> OutboundNumberPage:
    where: list[str] = []
    params: dict[str, Any] = {}
    if status in (0, 1):
        where.append("n.status = :status")
        params["status"] = status
    if keyword:
        where.append("n.phone LIKE :keyword")
        params["keyword"] = f"%{keyword}%"
    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    total = int(
        (
            await db.execute(
                text(f"SELECT COUNT(*) FROM admin_outbound_number n {where_sql}"), params
            )
        ).scalar()
        or 0
    )
    params["limit"] = page_size
    params["offset"] = (page - 1) * page_size
    rows = (
        await db.execute(
            text(f"{_NUMBER_SELECT} {where_sql} ORDER BY n.id DESC LIMIT :limit OFFSET :offset"),
            params,
        )
    ).mappings().all()
    return OutboundNumberPage(
        items=[_number_item(r) for r in rows],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


async def _assert_matchmaker(db: AsyncSession, matchmaker_id: int) -> None:
    ok = (await db.execute(text(_MATCHMAKER_VALID_SQL), {"mid": matchmaker_id})).scalar_one_or_none()
    if not ok:
        raise HTTPException(422, detail="绑定的红娘不存在或不是有效服务红娘")


async def create_number(
    db: AsyncSession, actor_id: int, body: OutboundNumberCreate
) -> OutboundNumberItem:
    if body.bound_matchmaker_id is not None:
        await _assert_matchmaker(db, body.bound_matchmaker_id)
    exists = (
        await db.execute(
            text("SELECT 1 FROM admin_outbound_number WHERE phone = :phone LIMIT 1"),
            {"phone": body.phone},
        )
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(409, detail="该外显号码已存在")
    result = await db.execute(
        text(
            "INSERT INTO admin_outbound_number (phone, area, bound_matchmaker_id, remark, created_by)"
            " VALUES (:phone, :area, :bound_matchmaker_id, :remark, :created_by)"
        ),
        {
            "phone": body.phone,
            "area": body.area,
            "bound_matchmaker_id": body.bound_matchmaker_id,
            "remark": body.remark,
            "created_by": actor_id,
        },
    )
    item = await _get_number(db, int(result.lastrowid))
    await _audit(db, actor_id, "outbound_number.create", "admin_outbound_number", item.id)
    await db.commit()
    return item


async def update_number(
    db: AsyncSession, actor_id: int, number_id: int, body: OutboundNumberUpdate
) -> OutboundNumberItem:
    await _get_number(db, number_id)
    values = body.model_dump(exclude_unset=True)
    if values.get("bound_matchmaker_id") is not None:
        await _assert_matchmaker(db, int(values["bound_matchmaker_id"]))
    if values:
        assignments = ", ".join(f"{key} = :{key}" for key in values)
        await db.execute(
            text(
                f"UPDATE admin_outbound_number SET {assignments}, updated_at = UTC_TIMESTAMP()"
                " WHERE id = :id"
            ),
            {**values, "id": number_id},
        )
    await _audit(db, actor_id, "outbound_number.update", "admin_outbound_number", number_id)
    await db.commit()
    return await _get_number(db, number_id)


async def delete_number(db: AsyncSession, actor_id: int, number_id: int) -> None:
    await _get_number(db, number_id)
    await db.execute(text("DELETE FROM admin_outbound_number WHERE id = :id"), {"id": number_id})
    await _audit(db, actor_id, "outbound_number.delete", "admin_outbound_number", number_id)
    await db.commit()


# ─────────────────────────── 外呼名单 ───────────────────────────


def _task_item(row: Any) -> OutboundTaskItem:
    data = dict(row)
    data["phone"] = mask_phone(data.get("phone")) or ""
    return OutboundTaskItem(**data)


async def _get_task(db: AsyncSession, task_id: int) -> OutboundTaskItem:
    row = (
        await db.execute(text(f"{_TASK_SELECT} WHERE t.id = :id"), {"id": task_id})
    ).mappings().first()
    if not row:
        raise HTTPException(404, detail="外呼名单不存在")
    return _task_item(row)


async def list_tasks(
    db: AsyncSession,
    page: int,
    page_size: int,
    status: str | None,
    matchmaker_id: int | None,
    keyword: str | None,
) -> OutboundTaskPage:
    where: list[str] = []
    params: dict[str, Any] = {}
    if status in {"pending", "calling", "done", "failed", "canceled"}:
        where.append("t.status = :status")
        params["status"] = status
    if matchmaker_id is not None:
        where.append("t.matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if keyword:
        where.append("(t.phone LIKE :keyword OR t.name LIKE :keyword)")
        params["keyword"] = f"%{keyword}%"
    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    total = int(
        (
            await db.execute(
                text(
                    "SELECT COUNT(*) FROM admin_outbound_task t LEFT JOIN users u"
                    " ON u.id = t.matchmaker_id " + where_sql
                ),
                params,
            )
        ).scalar()
        or 0
    )
    params["limit"] = page_size
    params["offset"] = (page - 1) * page_size
    rows = (
        await db.execute(
            text(f"{_TASK_SELECT} {where_sql} ORDER BY t.id DESC LIMIT :limit OFFSET :offset"),
            params,
        )
    ).mappings().all()
    return OutboundTaskPage(
        items=[_task_item(r) for r in rows],
        page=page,
        page_size=page_size,
        total=total,
        has_more=page * page_size < total,
    )


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    text_value = str(value).strip()
    return "" if text_value.lower() in {"nan", "none", "null"} else text_value


def _parse_import(data: bytes, filename: str) -> list[dict[str, str]]:
    lowered = (filename or "").lower()
    if lowered.endswith(".xls"):
        raise HTTPException(400, detail="暂不支持 .xls 旧格式，请用 Excel 另存为 .xlsx 后重新上传")
    if not lowered.endswith(".xlsx"):
        raise HTTPException(400, detail="仅支持 .xlsx 格式的 Excel 文件")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, detail=f"Excel 解析失败：{exc}") from exc
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [_cell_text(cell) for cell in (header_row or ())]

        def find(*keywords: str) -> int | None:
            for index, header in enumerate(headers):
                if any(keyword in header for keyword in keywords):
                    return index
            return None

        phone_col = find("手机", "电话", "phone")
        if phone_col is None:
            raise HTTPException(400, detail="模板表头缺少「手机号」列，请使用下载的模板填写")
        name_col = find("姓名", "昵称", "name")
        matchmaker_col = find("红娘", "跟进人", "matchmaker")
        rows: list[dict[str, str]] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if len(rows) >= _MAX_IMPORT_ROWS:
                raise HTTPException(
                    400, detail=f"单次最多导入 {_MAX_IMPORT_ROWS} 行，请拆分文件后重试"
                )
            cells = list(values or ())
            phone = _cell_text(cells[phone_col]) if phone_col < len(cells) else ""
            if not phone:
                continue
            name = (
                _cell_text(cells[name_col])
                if name_col is not None and name_col < len(cells)
                else ""
            )
            hint = (
                _cell_text(cells[matchmaker_col])
                if matchmaker_col is not None and matchmaker_col < len(cells)
                else ""
            )
            rows.append({"phone": phone, "name": name, "matchmaker": hint})
        return rows
    finally:
        workbook.close()


async def import_tasks(
    db: AsyncSession,
    actor_id: int,
    *,
    data: bytes,
    filename: str,
    default_matchmaker_id: int | None,
) -> OutboundTaskImportResult:
    rows = _parse_import(data, filename)
    if not rows:
        raise HTTPException(400, detail="文件内没有可导入的数据行")
    if default_matchmaker_id is not None:
        await _assert_matchmaker(db, default_matchmaker_id)
    batch_no = uuid.uuid4().hex[:24]
    created = 0
    errors: list[dict] = []
    matchmaker_cache: dict[str, int | None] = {}
    for index, row in enumerate(rows, start=2):
        matchmaker_id = default_matchmaker_id
        hint = row["matchmaker"]
        if hint:
            if hint not in matchmaker_cache:
                found = (
                    await db.execute(
                        text(
                            "SELECT id FROM users WHERE (nickname = :n OR real_name = :n)"
                            " AND status = 1 LIMIT 1"
                        ),
                        {"n": hint},
                    )
                ).scalar_one_or_none()
                matchmaker_cache[hint] = int(found) if found else None
            resolved = matchmaker_cache[hint]
            if resolved is None:
                errors.append({"row": index, "reason": f"红娘「{hint}」未找到"})
                continue
            matchmaker_id = resolved
        member_id = (
            await db.execute(
                text("SELECT id FROM users WHERE phone = :phone LIMIT 1"),
                {"phone": row["phone"]},
            )
        ).scalar_one_or_none()
        await db.execute(
            text(
                "INSERT INTO admin_outbound_task"
                " (batch_no, member_id, name, phone, matchmaker_id, status, created_by)"
                " VALUES (:batch_no, :member_id, :name, :phone, :matchmaker_id, 'pending', :created_by)"
            ),
            {
                "batch_no": batch_no,
                "member_id": int(member_id) if member_id else None,
                "name": row["name"] or None,
                "phone": row["phone"],
                "matchmaker_id": matchmaker_id,
                "created_by": actor_id,
            },
        )
        created += 1
    await _audit(
        db,
        actor_id,
        "outbound_task.import",
        "admin_outbound_task",
        None,
        reason=f"批次 {batch_no} 导入 {created} 条，跳过 {len(errors)} 条",
    )
    await db.commit()
    return OutboundTaskImportResult(
        batch_no=batch_no,
        total=len(rows),
        created=created,
        skipped=len(errors),
        errors=errors,
    )


async def call_task(
    db: AsyncSession, actor_id: int, task_id: int, caller_number: str | None
) -> OutboundTaskItem:
    """对名单行发起外呼。供应商未配置时在写库前 503。"""
    provider = get_outbound_provider()
    task = await _get_task(db, task_id)
    if task.status in {"done", "canceled"}:
        raise HTTPException(409, detail="该名单已完成或已取消，不能重复外呼")
    if caller_number:
        caller = (
            await db.execute(
                text(
                    "SELECT phone FROM admin_outbound_number"
                    " WHERE phone = :p AND status = 1 LIMIT 1"
                ),
                {"p": caller_number},
            )
        ).scalar_one_or_none()
        if not caller:
            raise HTTPException(422, detail="指定的外显号码不在号码池中或已停用")
    else:
        caller = (
            await db.execute(
                text(
                    "SELECT phone FROM admin_outbound_number WHERE status = 1"
                    " AND (bound_matchmaker_id = :mid OR bound_matchmaker_id IS NULL)"
                    " ORDER BY bound_matchmaker_id IS NULL, id LIMIT 1"
                ),
                {"mid": task.matchmaker_id},
            )
        ).scalar_one_or_none()

    # 真实号码取自库内行：列表回显已脱敏，不能拿脱敏值拨号。
    real_phone = (
        await db.execute(
            text("SELECT phone FROM admin_outbound_task WHERE id = :id"), {"id": task_id}
        )
    ).scalar_one()
    try:
        dialed = await provider.dial(real_phone, caller_number=caller, task_id=task_id)
    except HTTPException as exc:
        if exc.status_code == 503:
            # 供应商不可用：不改名单状态，保持 pending，配置后可直接重呼。
            raise
        await _mark_call_failed(db, actor_id, task_id, str(exc.detail))
        raise
    except Exception as exc:
        await _mark_call_failed(db, actor_id, task_id, str(exc))
        raise HTTPException(502, detail="外呼平台调用失败") from exc

    now = datetime.utcnow()
    await db.execute(
        text(
            "UPDATE admin_outbound_task SET status = 'done', result = :result,"
            " call_count = call_count + 1, last_called_at = :now,"
            " provider_call_id = :ref WHERE id = :id"
        ),
        {"result": dialed.result, "now": now, "ref": dialed.provider_call_id, "id": task_id},
    )
    await db.execute(
        text(
            "INSERT INTO admin_outbound_record (task_id, phone, matchmaker_id, direction, result,"
            " duration, recording_url, provider_call_id, started_at, ended_at)"
            " VALUES (:task_id, :phone, :matchmaker_id, 'outbound', :result, :duration,"
            " :recording_url, :ref, :now, :now)"
        ),
        {
            "task_id": task_id,
            "phone": real_phone,
            "matchmaker_id": task.matchmaker_id,
            "result": dialed.result,
            "duration": dialed.duration,
            "recording_url": dialed.recording_url,
            "ref": dialed.provider_call_id,
            "now": now,
        },
    )
    await _audit(db, actor_id, "outbound_task.call", "admin_outbound_task", task_id)
    await db.commit()
    return await _get_task(db, task_id)


async def _mark_call_failed(
    db: AsyncSession, actor_id: int, task_id: int, detail: str
) -> None:
    await db.execute(
        text(
            "UPDATE admin_outbound_task SET status = 'failed', call_count = call_count + 1,"
            " last_called_at = UTC_TIMESTAMP(), remark = :remark WHERE id = :id"
        ),
        {"remark": f"外呼失败：{detail[:200]}", "id": task_id},
    )
    await _audit(
        db, actor_id, "outbound_task.call", "admin_outbound_task", task_id, reason="外呼失败"
    )
    await db.commit()


# ─────────────────────────── 通话记录 ───────────────────────────


async def list_records(
    db: AsyncSession,
    page: int,
    page_size: int,
    result: str | None,
    matchmaker_id: int | None,
    start_date: str | None,
    end_date: str | None,
) -> OutboundRecordPage:
    where: list[str] = []
    params: dict[str, Any] = {}
    if result:
        where.append("r.result = :result")
        params["result"] = result
    if matchmaker_id is not None:
        where.append("r.matchmaker_id = :matchmaker_id")
        params["matchmaker_id"] = matchmaker_id
    if start_date:
        where.append("r.created_at >= :start_date")
        params["start_date"] = start_date
    if end_date:
        where.append("r.created_at < DATE_ADD(:end_date, INTERVAL 1 DAY)")
        params["end_date"] = end_date
    where_sql = f"WHERE {' AND '.join(where)}" if where else ""
    total = int(
        (
            await db.execute(
                text(
                    "SELECT COUNT(*) FROM admin_outbound_record r LEFT JOIN users u"
                    " ON u.id = r.matchmaker_id " + where_sql
                ),
                params,
            )
        ).scalar()
        or 0
    )
    params["limit"] = page_size
    params["offset"] = (page - 1) * page_size
    rows = (
        await db.execute(
            text(f"{_RECORD_SELECT} {where_sql} ORDER BY r.id DESC LIMIT :limit OFFSET :offset"),
            params,
        )
    ).mappings().all()
    items = []
    for row in rows:
        data = dict(row)
        data["phone"] = mask_phone(data.get("phone")) or ""
        data["duration"] = int(data.get("duration") or 0)
        items.append(OutboundRecordItem(**data))
    return OutboundRecordPage(
        items=items, page=page, page_size=page_size, total=total, has_more=page * page_size < total
    )


async def _audit(
    db: AsyncSession,
    actor_id: int,
    action: str,
    resource_type: str,
    resource_id: int | None,
    reason: str | None = None,
) -> None:
    await db.execute(
        text(
            "INSERT INTO business_audit_log"
            " (actor_user_id, action, resource_type, resource_id, reason)"
            " VALUES (:actor, :action, :resource_type, :resource_id, :reason)"
        ),
        {
            "actor": actor_id,
            "action": action,
            "resource_type": resource_type,
            "resource_id": resource_id,
            "reason": reason,
        },
    )
